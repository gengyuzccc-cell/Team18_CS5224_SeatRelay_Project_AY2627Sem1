#!/usr/bin/env python3
"""PowerPoint copy of the slides, for presenting from PowerPoint or Keynote.

    python3 -I report/tools/make_pptx.py                         # report/Slides.pdf -> report/Slides.pptx
    python3 -I report/tools/make_pptx.py --pdf report/Slides-submission.pdf --out report/Slides-submission.pptx

The Beamer PDF (report/Slides.tex) stays the source and the file that is submitted. Each PDF page
becomes one 16:9 slide holding a 200 dpi image of that page, so the PowerPoint copy looks exactly
like the PDF (its text is not editable; edit Slides.tex and re-run). The speaker notes name the
speaker printed in the slide footer; on the first slide of each speaker's part they carry that
part's "What to say" script from docs/walkthrough/part-N.md.
"""
from __future__ import annotations

import argparse
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches

ROOT = Path(__file__).resolve().parents[2]
WALK = ROOT / "docs" / "walkthrough"
DPI = 200


def page_texts(pdf: Path) -> list[str]:
    pages = int(re.search(r"Pages:\s+(\d+)", subprocess.run(
        ["pdfinfo", str(pdf)], check=True, capture_output=True, text=True).stdout).group(1))
    return [subprocess.run(["pdftotext", "-f", str(p), "-l", str(p), "-layout", str(pdf), "-"],
                           check=True, capture_output=True, text=True).stdout for p in range(1, pages + 1)]


def speaker_parts() -> dict[str, tuple[Path, str]]:
    """Speaker name -> (part file, its "What to say" section)."""
    parts = {}
    for path in sorted(WALK.glob("part-*.md")):
        text = path.read_text(encoding="utf-8")
        script = re.search(r"^## What to say.*?\n(.*?)(?=^## )", text, re.S | re.M)
        speaker = re.search(r"\*\*Speakers?:\*\*([^\n]*)", text)
        for name in re.findall(r"\b([A-Z]{2,}(?: [A-Z]{2,})+)\b", speaker.group(1) if speaker else ""):
            parts.setdefault(name, (path, script.group(1).strip() if script else ""))
    return parts


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--pdf", default=str(ROOT / "report" / "Slides.pdf"))
    ap.add_argument("--out", default=None, help="default: the PDF's name with .pptx")
    args = ap.parse_args()
    pdf = Path(args.pdf)
    out = Path(args.out) if args.out else pdf.with_suffix(".pptx")
    texts = page_texts(pdf)
    parts = speaker_parts()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    seen: set[str] = set()
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["pdftoppm", "-r", str(DPI), "-png", str(pdf), f"{tmp}/s"], check=True)
        images = sorted(Path(tmp).glob("s-*.png"))
        if len(images) != len(texts):
            sys.exit(f"{len(images)} images for {len(texts)} pages")
        for image, text in zip(images, texts):
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(image), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
            m = re.search(r"Speaker:\s+([A-Z]{2,}(?: [A-Z]{2,})+)", text)
            title = next((line.strip() for line in text.splitlines() if line.strip()), "")
            notes = [title]
            if m:
                name = m.group(1)
                part = parts.get(name)
                notes.append(f"Speaker: {name}" + (f" (script: docs/walkthrough/{part[0].name})" if part else ""))
                if part and name not in seen and not re.search(r"Backup B\d", text):
                    seen.add(name)
                    notes += ["", part[1]]
            slide.notes_slide.notes_text_frame.text = "\n".join(notes)
    prs.save(out)
    print(f"wrote {out.relative_to(ROOT) if out.is_relative_to(ROOT) else out} ({len(texts)} slides, {DPI} dpi)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
